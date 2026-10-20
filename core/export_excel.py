# -*- coding: utf-8 -*-
"""Xuất Excel tổng hợp – khổ A4 ngang, letterhead TEXO, bảng có lọc, cố định tiêu đề."""
import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .rules import GROUP_ORDER, GROUP_SHORT

ORG = "CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ"
DEPT = "PHÒNG KỸ THUẬT"
SIGNER = "Trưởng phòng Kỹ thuật"

NAVY = "1F3864"
GOLD = "C9A227"
LIGHT = "EEF2F8"
FONT = "Arial"

thin = Side(style="thin", color="B7C3D6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _page_setup(ws, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.oddHeader.left.text = "TEXO – " + DEPT.title()
    ws.oddHeader.left.size = 8
    ws.oddHeader.right.text = "Rà soát căn cứ pháp lý & tiêu chuẩn"
    ws.oddHeader.right.size = 8
    ws.oddFooter.left.text = "TEXO Legal Reviewer v2 – " + SIGNER
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Trang &P/&N"
    ws.oddFooter.right.size = 8
    if title_rows:
        ws.print_title_rows = title_rows


def _letterhead(ws, title, ncols, subtitle=""):
    last = get_column_letter(max(ncols, 4))
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = f"{ORG}  |  {DEPT}"
    ws["A1"].font = Font(name=FONT, size=9, bold=True, color=NAVY)
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    ws["A2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 26
    ws.merge_cells(f"A3:{last}3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name=FONT, size=9, italic=True, color="555555")
    for c in range(1, max(ncols, 4) + 1):
        ws.cell(row=2, column=c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(row=3, column=c).border = Border(bottom=Side(style="medium", color=GOLD))


def _write_table(ws, df, start_row, widths=None, wrap_cols=(), link_col=None, center_cols=()):
    cols = list(df.columns)
    for j, col in enumerate(cols, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[start_row].height = 32
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, val in enumerate(row, 1):
            col = cols[j - 1]
            try:
                if val is None or pd.isna(val):
                    val = None
            except (TypeError, ValueError):
                pass
            if hasattr(val, "item"):
                val = val.item()
            cell = ws.cell(row=i, column=j)
            if col == link_col and val:
                cell.value = "🔎 Tra cứu"
                cell.hyperlink = str(val)
                cell.font = Font(name=FONT, size=10, color="1F5FBF", underline="single")
            else:
                cell.value = val
                cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=col in wrap_cols,
                                       horizontal="center" if col in center_cols else None)
            cell.border = BORDER
            if (i - start_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    if widths:
        for j, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 14)
    end_row = start_row + max(len(df), 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{end_row}"
    freeze_col = (cols.index(link_col) + 2) if link_col in cols else (3 if len(cols) > 3 else 1)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=freeze_col)
    return end_row


def build_excel(file_items, items, sum_df, cat_df, mat_df, det_df, mode_text, warnings):
    wb = Workbook()
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    sub = f"Ngày lập: {now}   •   Số file: {len(file_items)}   •   Chế độ: {mode_text}"

    # ---------------- Sheet 1: Tổng hợp
    ws = wb.active
    ws.title = "Tổng hợp"
    ncol = max(len(sum_df.columns), 8)
    _letterhead(ws, "BÁO CÁO TRÍCH XUẤT CĂN CỨ PHÁP LÝ & TIÊU CHUẨN", ncol, sub)
    # KPI
    ws["A5"] = "TỔNG QUAN"
    ws["A5"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    kpis = [("Tổng số văn bản", len(items))] + \
           [(GROUP_SHORT[g], sum(1 for c in items if c["group"] == g)) for g in GROUP_ORDER] + \
           [("Mục cần lưu ý", sum(1 for c in items if c["flags"]))]
    for j, (lab, val) in enumerate(kpis, 1):
        a = ws.cell(row=6, column=j, value=lab)
        a.font = Font(name=FONT, size=9, color="555555")
        a.alignment = Alignment(horizontal="center", wrap_text=True)
        b = ws.cell(row=7, column=j, value=val)
        b.font = Font(name=FONT, size=18, bold=True, color=NAVY if lab != "Mục cần lưu ý" else "C00000")
        b.alignment = Alignment(horizontal="center")
        for r in (6, 7):
            ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor=LIGHT)
            ws.cell(row=r, column=j).border = BORDER
    ws.row_dimensions[6].height = 28
    ws.row_dimensions[7].height = 30
    ws["A9"] = "THEO TỪNG FILE"
    ws["A9"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
    widths = {"Tên file": 42, "Định dạng": 10, "Số trang": 9, "Tổng số văn bản": 11, "Cảnh báo": 60}
    end = _write_table(ws, sum_df, 10, widths, wrap_cols=("Tên file", "Cảnh báo"),
                       center_cols=tuple(c for c in sum_df.columns if c not in ("Tên file", "Cảnh báo")))
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    r = end + 2
    if warnings:
        ws.cell(row=r, column=1, value="CẢNH BÁO KHI XỬ LÝ").font = Font(name=FONT, size=11, bold=True, color="C00000")
        for w in warnings:
            r += 1
            ws.cell(row=r, column=1, value="• " + w).font = Font(name=FONT, size=9)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Lưu ý: Danh mục trích xuất tự động để hỗ trợ rà soát. Tình trạng hiệu lực phải "
                                   "tra cứu tại nguồn chính thức (link 🔎 trong sheet Danh mục). Nhận xét của AI "
                                   "(nếu có) chỉ để tham khảo.").font = Font(name=FONT, size=9, italic=True, color="555555")
    _page_setup(ws)

    # ---------------- Sheet 2: Danh mục
    ws = wb.create_sheet("Danh mục văn bản")
    _letterhead(ws, "DANH MỤC VĂN BẢN PHÁP LUẬT & TIÊU CHUẨN ĐƯỢC TRÍCH DẪN", len(cat_df.columns), sub)
    widths = {"STT": 6, "Nhóm": 18, "Loại": 20, "Số hiệu / Mã hiệu": 30, "Tên – trích yếu (theo hồ sơ)": 46,
              "Ngày ban hành": 12, "Số file": 7, "Số lần trích": 8, "File trích dẫn": 34,
              "Ghi chú kiểm tra": 40, "Nguồn": 10, "Hiệu lực (AI)": 18, "Số hiệu chuẩn (AI)": 22,
              "Tên đầy đủ (AI)": 40, "VB thay thế/sửa đổi (AI)": 22, "Nhận xét AI": 44, "Tra cứu hiệu lực": 12}
    wrap = ("Tên – trích yếu (theo hồ sơ)", "File trích dẫn", "Ghi chú kiểm tra", "Tên đầy đủ (AI)",
            "Nhận xét AI", "Số hiệu / Mã hiệu", "Nhóm", "Loại")
    _write_table(ws, cat_df, 5, widths, wrap_cols=wrap, link_col="Tra cứu hiệu lực",
                 center_cols=("STT", "Ngày ban hành", "Số file", "Số lần trích", "Nguồn"))
    # tô đỏ nhẹ dòng có ghi chú kiểm tra
    if "Ghi chú kiểm tra" in cat_df.columns:
        jc = list(cat_df.columns).index("Ghi chú kiểm tra") + 1
        for i in range(len(cat_df)):
            cell = ws.cell(row=6 + i, column=jc)
            if cell.value:
                cell.font = Font(name=FONT, size=10, color="C00000")
    _page_setup(ws, "5:5")

    # ---------------- Sheet 3: Ma trận
    ws = wb.create_sheet("Ma trận VB × File")
    _letterhead(ws, "MA TRẬN VĂN BẢN × FILE (số lần trích dẫn)", len(mat_df.columns), sub)
    widths = {"STT": 6, "Số hiệu / Mã hiệu": 34, "Tra cứu hiệu lực": 12, "Nhóm": 13, "Tổng": 8}
    for c in mat_df.columns[4:-1]:
        widths[c] = 16
    _write_table(ws, mat_df, 5, widths, wrap_cols=("Số hiệu / Mã hiệu",), link_col="Tra cứu hiệu lực",
                 center_cols=tuple(mat_df.columns[3:]))
    for j in range(5, len(mat_df.columns)):
        ws.cell(row=5, column=j).alignment = Alignment(horizontal="center", vertical="center",
                                                       wrap_text=True, text_rotation=0)
    ws.row_dimensions[5].height = 60
    _page_setup(ws, "5:5")

    # ---------------- Sheet 4: Chi tiết
    ws = wb.create_sheet("Chi tiết trích dẫn")
    _letterhead(ws, "CHI TIẾT VỊ TRÍ TRÍCH DẪN", max(len(det_df.columns), 4), sub)
    widths = {"STT": 6, "File": 30, "Vị trí": 16, "Văn bản": 30, "Tra cứu hiệu lực": 12, "Nhóm": 12, "Nguyên văn": 30,
              "Ngữ cảnh": 70, "Nguồn": 16}
    if not det_df.empty:
        _write_table(ws, det_df, 5, widths, wrap_cols=("File", "Văn bản", "Nguyên văn", "Ngữ cảnh"),
                     link_col="Tra cứu hiệu lực",
                     center_cols=("STT",))
    _page_setup(ws, "5:5")

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
